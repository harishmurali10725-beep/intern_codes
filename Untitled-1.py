from openpyxl import Workbook
import json
wb = Workbook()
sheet = wb.active
input = None
with open('input.json','r') as file:
    input = json.load(file)
input = input["inputs"]
columns = list(set([col for row in range(len(input)) for col in input[row].keys()]))
print(columns)
for i in range(len(columns)):
    sheet.cell(row=1,column=i+1).value = columns[i]
for r, row_data in enumerate(input, start=2):
    for c, col in enumerate(columns, start=1):
        sheet.cell(row=r, column=c).value = row_data.get(col, " ")

wb.save("output1.xlsx")
